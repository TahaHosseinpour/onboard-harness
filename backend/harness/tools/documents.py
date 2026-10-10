"""Document tools: read Word, PowerPoint and PDF files as text, and write Word and PowerPoint files.

`read_file` only understands plain text (and shows images and PDFs to models
that can look at them). Office files are zipped XML, so they need code to get
the words out - and that also makes PDFs readable for models that cannot see.

Writing goes the other way: the agent describes the document in a small subset
of Markdown and code lays it out. Persian and Arabic text is set right-to-left.
"""

from __future__ import annotations

import re
from pathlib import Path

from langchain_core.tools import BaseTool, tool

from ..workspace import Workspace
from .legacy import read_doc, read_ppt
from .paths import existing_file, new_file

_RTL = re.compile(r"[֐-ࣿיִ-﷿ﹰ-﻿]")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_NUMBERED = re.compile(r"^\d+[.)]\s+")


def _row(cells: list[str]) -> str:
    return "| " + " | ".join(cell.replace("\n", " ").replace("|", "/").strip() for cell in cells) + " |"


def _table(rows: list[list[str]]) -> list[str]:
    """A table as Markdown lines: readable as plain text, and drawn as a table where Markdown is rendered."""
    if not rows:
        return []
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    return ["", _row(padded[0]), "|" + " --- |" * width, *(_row(row) for row in padded[1:]), ""]


# ── reading ──


def _read_docx(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    lines: list[str] = []
    # Paragraphs and tables in the order they appear on the page.
    for block in document.element.body.iterchildren():
        if block.tag.endswith("}p"):
            paragraph = Paragraph(block, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "") if paragraph.style is not None else ""
            if style.startswith("Heading"):
                level = int(style[-1]) if style[-1].isdigit() else 1
                text = "#" * level + " " + text
            elif style.startswith("List"):
                text = "- " + text
            lines.append(text)
        elif block.tag.endswith("}tbl"):
            lines.extend(_table([[cell.text for cell in row.cells] for row in Table(block, document).rows]))
    return "\n".join(lines)


def _read_pptx(path: Path) -> str:
    from pptx import Presentation

    lines: list[str] = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        lines.append(f"--- slide {number} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                lines.extend(_table([[cell.text for cell in row.cells] for row in shape.table.rows]))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"(speaker notes) {notes}")
    return "\n".join(lines)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("this PDF is protected with a password")
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    if not any(pages):
        raise ValueError("this PDF has no text layer (it is probably a scan). Look at it with read_file instead, if the model can see images")
    return "\n".join(f"--- page {number} ---\n{text}" for number, text in enumerate(pages, start=1))


_READERS = {".docx": _read_docx, ".pptx": _read_pptx, ".pdf": _read_pdf, ".doc": read_doc, ".ppt": read_ppt}


def document_markdown(path: Path) -> str:
    """A Word or PowerPoint file as Markdown, for showing it in the app. Raises ValueError for other files."""
    suffix = path.suffix.lower()
    if suffix not in (".docx", ".pptx", ".doc", ".ppt"):
        raise ValueError("only Word and PowerPoint files can be shown as a document")
    text = _READERS[suffix](path)
    if suffix in (".doc", ".ppt", ".pptx"):
        text = _plain_to_markdown(text)
    # Slides are separated by a rule and numbered, so a deck reads as a sequence of pages.
    return re.sub(r"^--- slide (\d+) ---$", r"\n---\n\n###### اسلاید \1", text, flags=re.MULTILINE).strip()


def _plain_to_markdown(text: str) -> str:
    """Plain extracted text as Markdown: every line its own paragraph, runs of `a | b` lines as tables."""
    out: list[str] = []
    table: list[list[str]] = []
    for line in [*text.splitlines(), ""]:
        if " | " in line and not line.startswith("|"):
            table.append(line.split(" | "))
            continue
        if table:
            out.extend(_table(table))
            table = []
        out.extend([line, ""])
    return "\n".join(out)


# ── writing ──


def _is_rtl(text: str) -> bool:
    return bool(_RTL.search(text))


def _table_rows(lines: list[str]) -> list[list[str]]:
    """Markdown table lines as rows of cells, without the `|---|` separator row."""
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    return [row for row in rows if not all(set(cell) <= set("-: ") for cell in row)]


def _blocks(content: str) -> list[tuple[str, object]]:
    """Split the Markdown subset into (kind, value): heading, bullet, number, table, text."""
    blocks: list[tuple[str, object]] = []
    table: list[str] = []
    for raw in content.splitlines() + [""]:
        line = raw.strip()
        if line.startswith("|"):
            table.append(line)
            continue
        if table:
            blocks.append(("table", _table_rows(table)))
            table = []
        if not line:
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append((f"heading{min(level, 3)}", line.lstrip("#").strip()))
        elif line[:2] in ("- ", "* "):
            blocks.append(("bullet", line[2:].strip()))
        elif _NUMBERED.match(line):
            blocks.append(("number", _NUMBERED.sub("", line)))
        else:
            blocks.append(("text", line))
    return blocks


def _write_docx(path: Path, content: str) -> int:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()

    def fill(paragraph, text: str) -> None:
        rtl = _is_rtl(text)
        if rtl:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
        # `**bold**` is the one inline mark supported.
        for index, part in enumerate(_BOLD.split(text)):
            if not part:
                continue
            run = paragraph.add_run(part)
            run.bold = bool(index % 2) or None
            if rtl:
                run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))

    styles = {"bullet": "List Bullet", "number": "List Number", "text": None}
    for kind, value in _blocks(content):
        if kind == "table":
            rows: list[list[str]] = value  # type: ignore[assignment]
            if not rows:
                continue
            table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
            table.style = "Table Grid"
            if any(_is_rtl(cell) for row in rows for cell in row):
                properties = table._tbl.tblPr
                visual = OxmlElement("w:bidiVisual")
                visual.set(qn("w:val"), "1")
                properties.append(visual)
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    fill(table.cell(r, c).paragraphs[0], cell)
                    if r == 0:
                        for run in table.cell(r, c).paragraphs[0].runs:
                            run.bold = True
        elif kind.startswith("heading"):
            fill(document.add_heading(level=int(kind[-1])), str(value))
        else:
            fill(document.add_paragraph(style=styles[kind]), str(value))
    document.save(str(path))
    return len(document.paragraphs)


def _write_pptx(path: Path, content: str) -> int:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    deck = Presentation()
    layout = deck.slide_layouts[1]  # title and content

    def fill(paragraph, text: str, size: int | None = None) -> None:
        paragraph.text = _BOLD.sub(r"\1", text)
        if size:
            for run in paragraph.runs:
                run.font.size = Pt(size)
        if _is_rtl(text):
            paragraph.alignment = PP_ALIGN.RIGHT
            paragraph._p.get_or_add_pPr().set("rtl", "1")

    body = None
    first_line = True
    for kind, value in _blocks(content):
        if kind == "heading1" or body is None:
            slide = deck.slides.add_slide(layout)
            fill(slide.shapes.title.text_frame.paragraphs[0], str(value) if kind.startswith("heading") else "")
            body, first_line = slide.placeholders[1].text_frame, True
            if kind.startswith("heading"):
                continue
        lines = [_row(row) for row in value] if kind == "table" else [str(value)]  # type: ignore[union-attr]
        for line in lines:
            paragraph = body.paragraphs[0] if first_line else body.add_paragraph()
            first_line = False
            fill(paragraph, line, size=20)
    if not deck.slides:
        raise ValueError("no slides: start each slide with a line like `# Slide title`")
    deck.save(str(path))
    return len(deck.slides)


_WRITERS = {".docx": _write_docx, ".pptx": _write_pptx}


def make_document_tools(workspace: Workspace) -> list[BaseTool]:
    @tool
    def read_document(path: str, start: int = 0, max_chars: int = 20000) -> str:
        """Read the text of a Word (.docx, .doc), PowerPoint (.pptx, .ppt) or PDF file in the workspace.

        Use this instead of `read_file` for these formats. Headings, lists,
        tables, slide numbers and page numbers are kept as plain text (the old
        .doc and .ppt formats give plain text only). A long document comes in
        parts: call again with `start` set to where the last part ended. For
        .xlsx, .xls and .csv use `read_table`.
        """
        try:
            file = existing_file(workspace, path)
            suffix = file.suffix.lower()
            if suffix not in _READERS:
                return f"Error: read_document reads Word, PowerPoint and PDF files. For {suffix or 'this file'} use read_file (text, images) or read_table (.xlsx, .xls, .csv)."
            text = _READERS[suffix](file)
        except ValueError as error:
            return f"Error: {error}."
        except Exception as error:  # noqa: BLE001 - a damaged document is an answer for the agent, not a crash
            return f"Error: this file could not be read as {Path(path).suffix} ({type(error).__name__}). It may be damaged or password-protected."
        if not text.strip():
            return "The document has no text."
        start, max_chars = max(0, start), max(1000, min(max_chars, 60000))
        part = text[start : start + max_chars]
        rest = len(text) - start - len(part)
        return part + (f"\n\n[{rest} more characters. Call read_document again with start={start + len(part)}.]" if rest > 0 else "")

    @tool
    def write_document(path: str, content: str, overwrite: bool = False) -> str:
        """Create a Word (.docx) or PowerPoint (.pptx) file in the workspace from simple Markdown.

        Supported: `# Heading`, `## Sub-heading`, `- bullet`, `1. numbered`,
        `**bold**`, tables written as `| a | b |` rows, and plain paragraphs.
        For .pptx every `# Heading` starts a new slide and the lines under it
        become that slide's points; keep each slide to a few short lines.
        Persian and Arabic text is laid out right-to-left automatically.
        An existing file is not replaced unless `overwrite` is true.
        """
        suffix = Path(path).suffix.lower()
        if suffix not in _WRITERS:
            return "Error: write_document creates .docx and .pptx files (not the old .doc / .ppt). For plain text or Markdown use write_file; for spreadsheets use write_table."
        if not content.strip():
            return "Error: the content is empty."
        try:
            count = _WRITERS[suffix](new_file(workspace, path, overwrite=overwrite), content)
        except (ValueError, OSError) as error:
            return f"Error: {error}."
        return f"Created {path} ({count} {'slides' if suffix == '.pptx' else 'paragraphs'})."

    return [read_document, write_document]
