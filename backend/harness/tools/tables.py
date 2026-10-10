"""Table tools: read CSV / Excel files, total them with code, add rows to them, and create new ones."""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path
from typing import Any, Literal

from langchain_core.tools import BaseTool, tool

from ..workspace import Workspace
from .paths import existing_file, new_file

_OPERATIONS = {
    "sum": sum,
    "avg": lambda xs: sum(xs) / len(xs),
    "min": min,
    "max": max,
    "count": len,
}


def _sheet(book: Any, name: str | None) -> Any:
    """The named sheet of a workbook, or its first one."""
    if name is None:
        return book.worksheets[0]
    if name not in book.sheetnames:
        raise ValueError(f"no sheet {name!r}. Sheets are: {', '.join(book.sheetnames)}")
    return book[name]


def _xls_cell(book: Any, cell: Any) -> object:
    """A cell of an old .xls sheet as a plain value: dates as dates, whole numbers without `.0`."""
    import xlrd

    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            moment = xlrd.xldate_as_datetime(cell.value, book.datemode)
        except (ValueError, OverflowError):
            return cell.value
        return moment.date().isoformat() if moment.time().isoformat() == "00:00:00" else moment.isoformat(sep=" ")
    if cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
        return int(cell.value)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    return cell.value


def _load_xls(path: Path, sheet_name: str | None) -> list[dict[str, object]]:
    import xlrd

    try:
        book = xlrd.open_workbook(str(path))
    except xlrd.XLRDError as error:
        raise ValueError(f"this .xls file could not be read ({error})") from error
    if sheet_name is not None and sheet_name not in book.sheet_names():
        raise ValueError(f"no sheet {sheet_name!r}. Sheets are: {', '.join(book.sheet_names())}")
    sheet = book.sheet_by_name(sheet_name) if sheet_name else book.sheet_by_index(0)
    if sheet.nrows == 0:
        return []
    header = [str(_xls_cell(book, cell)) if cell.value not in ("", None) else f"column_{i + 1}" for i, cell in enumerate(sheet.row(0))]
    rows = ([_xls_cell(book, cell) for cell in sheet.row(index)] for index in range(1, sheet.nrows))
    return [dict(zip(header, row, strict=False)) for row in rows if any(value is not None for value in row)]


def _sheet_names(path: Path) -> list[str]:
    if path.suffix.lower() == ".xls":
        import xlrd

        try:
            return xlrd.open_workbook(str(path), on_demand=True).sheet_names()
        except xlrd.XLRDError as error:
            raise ValueError(f"this .xls file could not be read ({error})") from error
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return []
    from openpyxl import load_workbook

    return load_workbook(path, read_only=True).sheetnames


def _load(path: Path, sheet_name: str | None = None) -> list[dict[str, object]]:
    """Rows of a table as dicts keyed by the header row."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook  # imported lazily: only needed for Excel files

        sheet = _sheet(load_workbook(path, read_only=True, data_only=True), sheet_name)
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell) if cell is not None else f"column_{i + 1}" for i, cell in enumerate(next(rows, ()))]
        return [dict(zip(header, row, strict=False)) for row in rows if any(cell is not None for cell in row)]
    if path.suffix.lower() == ".xls":
        return _load_xls(path, sheet_name)
    raise ValueError("only .csv, .xlsx and .xls files are supported")


def _append(path: Path, rows: list[dict[str, Any]], sheet_name: str | None = None) -> list[str]:
    """Add rows under the table's existing header, and return that header. Nothing already in the file is touched."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            header = next(csv.reader(handle), [])
            ends_with_newline = handle.read()[-1:] in ("", "\n")
        _check_columns(header, rows)
        with path.open("a", encoding="utf-8", newline="") as handle:
            if not ends_with_newline:
                handle.write("\n")
            csv.DictWriter(handle, fieldnames=header, restval="").writerows(rows)
        return header
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        book = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
        sheet = _sheet(book, sheet_name)
        header = [str(cell.value) for cell in sheet[1] if cell.value is not None]
        _check_columns(header, rows)
        for row in rows:
            sheet.append([row.get(column) for column in header])
        book.save(path)
        return header
    if path.suffix.lower() == ".xls":
        raise ValueError("the old .xls format can be read but not changed. Copy the table to a new .xlsx with write_table and add the rows there")
    raise ValueError("only .csv and .xlsx files are supported")


def _check_columns(header: list[str], rows: list[dict[str, Any]]) -> None:
    if not header:
        raise ValueError("the table has no header row to match the new rows against")
    unknown = sorted({str(key) for row in rows for key in row} - set(header))
    if unknown:
        raise ValueError(f"no such columns: {', '.join(unknown)}. Columns are: {', '.join(header)}")


def table_preview(path: Path, sheet_name: str | None = None, limit: int = 500) -> dict[str, Any]:
    """The first rows of a .csv or .xlsx sheet as plain values, for showing the table in the app."""
    names = _sheet_names(path)
    rows = _load(path, sheet_name)
    columns = list(rows[0]) if rows else []
    cells = [["" if row.get(column) is None else row.get(column) for column in columns] for row in rows[:limit]]
    # Dates and other objects are shown as text; numbers stay numbers so the screen can align them.
    cells = [[value if isinstance(value, (int, float, str)) and not isinstance(value, bool) else str(value) for value in row] for row in cells]
    return {"sheets": names, "sheet": sheet_name or (names[0] if names else None), "columns": columns, "rows": cells, "total_rows": len(rows), "truncated": len(rows) > limit}


def _number(value: object) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    try:
        return float(str(value).replace(",", "").replace("٬", "").strip())
    except ValueError:
        return None


def _create(path: Path, columns: list[str], rows: list[list[Any]], sheet_name: str | None) -> None:
    if path.suffix.lower() == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as handle:  # the BOM makes Excel show Persian text correctly
            csv.writer(handle).writerows([columns, *rows])
        return
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font

    adding = path.exists() and bool(sheet_name)
    book = load_workbook(path) if adding else Workbook()
    if adding and sheet_name in book.sheetnames:
        raise ValueError(f"sheet {sheet_name!r} already exists in this file")
    sheet = book.create_sheet(sheet_name) if adding else book.active
    if sheet_name:
        sheet.title = sheet_name
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    book.save(path)


def make_table_tools(workspace: Workspace) -> list[BaseTool]:
    def resolve(path: str) -> Path:
        return existing_file(workspace, path)

    @tool
    def read_table(path: str, max_rows: int = 20, sheet: str | None = None) -> str:
        """Show the columns, row count and first rows of a .csv, .xlsx or .xls file in the workspace.

        Call this before analysing a table so you use its real column names.
        An Excel file can hold several sheets: their names are listed, and
        `sheet` picks one (default: the first).
        """
        try:
            file = resolve(path)
            names = _sheet_names(file)
            rows = _load(file, sheet)
        except ValueError as error:
            return f"Error: {error}."
        sheets = f"sheets: {', '.join(names)} (showing {sheet or names[0]})\n" if len(names) > 1 else ""
        if not rows:
            return sheets + "The table is empty."
        columns = list(rows[0])
        preview = "\n".join(" | ".join(str(row.get(column, "")) for column in columns) for row in rows[: max(1, max_rows)])
        return f"{sheets}columns: {', '.join(columns)}\nrows: {len(rows)}\n\n{' | '.join(columns)}\n{preview}"

    @tool
    def aggregate_table(
        path: str,
        value_column: str,
        operation: Literal["sum", "avg", "min", "max", "count"] = "sum",
        group_by: str | None = None,
        sheet: str | None = None,
    ) -> str:
        """Compute sum / avg / min / max / count of a numeric column, optionally per group.

        Use this for every total you report from a table. Example:
        aggregate_table("/sales.csv", value_column="amount", operation="sum", group_by="city")
        """
        try:
            rows = _load(resolve(path), sheet)
        except ValueError as error:
            return f"Error: {error}."
        if rows and value_column not in rows[0]:
            return f"Error: no column {value_column!r}. Columns are: {', '.join(rows[0])}."
        if group_by and rows and group_by not in rows[0]:
            return f"Error: no column {group_by!r}. Columns are: {', '.join(rows[0])}."

        groups: dict[str, list[float]] = defaultdict(list)
        skipped = 0
        for row in rows:
            number = _number(row.get(value_column))
            if number is None:
                skipped += 1
                continue
            groups[str(row.get(group_by)) if group_by else "all"].append(number)
        if not groups:
            return f"Error: column {value_column!r} has no numeric values."

        lines = [f"{key}: {round(_OPERATIONS[operation](values), 4)}" for key, values in sorted(groups.items())]
        note = f"\n({skipped} rows without a number in {value_column!r} were skipped)" if skipped else ""
        return f"{operation} of {value_column}" + (f" by {group_by}" if group_by else "") + ":\n" + "\n".join(lines) + note

    @tool
    def add_table_rows(path: str, rows: list[dict[str, Any]], sheet: str | None = None) -> str:
        """Add rows to the end of an existing .csv or .xlsx table in the workspace.

        Each row is an object keyed by the table's own column names (call
        `read_table` first to get them); a column you leave out stays empty.
        Existing rows, formulas and formatting are not changed. Example:
        add_table_rows("/expenses.xlsx", rows=[{"date": "1405/07/12", "title": "Taxi", "amount": 250000}])
        """
        if not rows:
            return "Error: no rows were given."
        try:
            header = _append(resolve(path), rows, sheet)
        except (ValueError, OSError) as error:
            return f"Error: {error}."
        return f"Added {len(rows)} rows to {path} (columns: {', '.join(header)})."

    @tool
    def write_table(path: str, columns: list[str], rows: list[list[Any]], sheet: str | None = None, overwrite: bool = False) -> str:
        """Create a new .xlsx or .csv table in the workspace.

        `columns` is the header row; each item of `rows` is one row, with values
        in the same order. Numbers given as numbers stay numbers in Excel. To
        add a sheet to an existing .xlsx file, give a new `sheet` name. An
        existing file is not replaced unless `overwrite` is true; to add rows
        to an existing table use `add_table_rows`.
        """
        suffix = Path(path).suffix.lower()
        if suffix not in {".xlsx", ".csv"}:
            return "Error: write_table creates .xlsx and .csv files."
        if not columns:
            return "Error: give the column names."
        try:
            adding_sheet = bool(sheet) and suffix == ".xlsx" and not overwrite
            file = new_file(workspace, path, overwrite=overwrite or adding_sheet)
            if overwrite and file.exists():
                file.unlink()
            _create(file, [str(column) for column in columns], rows, sheet)
        except (ValueError, OSError) as error:
            return f"Error: {error}."
        return f"Created {path}" + (f" (sheet {sheet})" if sheet else "") + f" with {len(rows)} rows."

    return [read_table, aggregate_table, add_table_rows, write_table]
