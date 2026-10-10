"""The file-format and browser tools, exercised on real files they make and read back."""

from __future__ import annotations

import http.server
import subprocess
import sys
import threading
import zipfile

import pytest

from harness.permissions import build_policy, load_rules
from harness.tools import browser, build_tools


@pytest.fixture
def tools(workspace):
    return build_tools(workspace)


def _pdf(text: str) -> bytes:
    """A one-page PDF with a line of text, built by hand so the test needs no PDF writer."""
    stream = f"BT /F1 18 Tf 40 700 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    body, offsets = b"%PDF-1.4\n", []
    for number, content in enumerate(objects, start=1):
        offsets.append(len(body))
        body += f"{number} 0 obj\n".encode() + content + b"\nendobj\n"
    xref = len(body)
    body += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode() + b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    return body + f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()


# ── Word, PowerPoint, PDF ──

REPORT = """# گزارش فروش مهر
فروش این ماه **۲٬۸۴۰ میلیون** تومان بود.

## Cities
- Tehran
- Shiraz

| city | amount |
|---|---|
| Tehran | 1200 |
| شیراز | 900 |
"""


def test_a_word_file_is_written_and_read_back(tools, workspace):
    from docx import Document

    assert tools["write_document"].invoke({"path": "/out/report.docx", "content": REPORT}).startswith("Created")
    text = tools["read_document"].invoke({"path": "/out/report.docx"})

    assert "# گزارش فروش مهر" in text and "## Cities" in text and "- Tehran" in text
    assert "Tehran | 1200" in text and "شیراز | 900" in text
    assert "**" not in text, "bold marks become formatting, not stars"
    first = Document(str(workspace.root / "out" / "report.docx")).paragraphs[0]
    assert first._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi") is not None, "Persian text is right-to-left"


def test_a_presentation_is_written_and_read_back(tools):
    deck = "# فروش مهر\n- رشد ۱۲ درصدی\n- Tehran leads\n\n# Next steps\n1. Hire\n2. Ship"
    assert tools["write_document"].invoke({"path": "/deck.pptx", "content": deck}) == "Created /deck.pptx (2 slides)."
    text = tools["read_document"].invoke({"path": "/deck.pptx"})

    assert text.index("--- slide 1 ---") < text.index("رشد ۱۲ درصدی") < text.index("--- slide 2 ---") < text.index("Ship")


def test_a_pdf_is_read_as_text_and_long_documents_come_in_parts(tools, workspace):
    (workspace.root / "invoice.pdf").write_bytes(_pdf("Invoice 1042 total 250000"))
    assert "Invoice 1042 total 250000" in tools["read_document"].invoke({"path": "/invoice.pdf"})

    tools["write_document"].invoke({"path": "/long.docx", "content": "\n\n".join(f"paragraph number {n} " + "x" * 80 for n in range(40))})
    first = tools["read_document"].invoke({"path": "/long.docx", "max_chars": 1000})
    assert "start=1000" in first and "paragraph number 39" not in first
    assert "paragraph number 39" in tools["read_document"].invoke({"path": "/long.docx", "start": 1000})


def test_documents_fail_with_a_reason_and_never_overwrite_by_accident(tools, workspace):
    (workspace.root / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0")
    (workspace.root / "broken.docx").write_bytes(b"not a zip")
    tools["write_document"].invoke({"path": "/a.docx", "content": "first"})

    assert tools["read_document"].invoke({"path": "/old.doc"}).startswith("Error:")
    assert tools["read_document"].invoke({"path": "/broken.docx"}).startswith("Error:")
    assert tools["read_document"].invoke({"path": "/../outside.docx"}).startswith("Error:")
    assert "already exists" in tools["write_document"].invoke({"path": "/a.docx", "content": "second"})
    assert tools["write_document"].invoke({"path": "/a.docx", "content": "second", "overwrite": True}).startswith("Created")
    assert tools["write_document"].invoke({"path": "/../outside.docx", "content": "x"}).startswith("Error:")


def test_old_office_files_are_read(tools, workspace):
    """The samples were saved by Microsoft Office itself in the 97-2003 formats."""
    import shutil
    from pathlib import Path

    for sample in (Path(__file__).parent / "samples").iterdir():
        shutil.copy(sample, workspace.root / sample.name)

    word = tools["read_document"].invoke({"path": "/sample.doc"})
    assert "Quarterly sales report" in word and "گزارش فروش فصل پاییز" in word and "Tehran | 1200" in word
    deck = tools["read_document"].invoke({"path": "/sample.ppt"})
    assert deck.index("--- slide 1 ---") < deck.index("رشد دوازده درصدی") < deck.index("--- slide 2 ---") < deck.index("Hire two people")
    assert "Click to edit" not in deck, "the master's placeholder prompts are not the deck's text"

    sheet = tools["read_table"].invoke({"path": "/sample.xls"})
    assert "columns: city, amount" in sheet and "Tehran | 1200" in sheet and "شیراز | 900.5" in sheet
    assert "all: 2100.5" in tools["aggregate_table"].invoke({"path": "/sample.xls", "value_column": "amount"})
    assert "can be read but not changed" in tools["add_table_rows"].invoke({"path": "/sample.xls", "rows": [{"city": "Yazd"}]})


# ── Excel ──


def test_a_workbook_is_created_with_sheets_and_read_per_sheet(tools):
    assert tools["write_table"].invoke({"path": "/book.xlsx", "columns": ["city", "amount"], "rows": [["Tehran", 1200], ["شیراز", 900]], "sheet": "مهر"}).startswith("Created")
    assert tools["write_table"].invoke({"path": "/book.xlsx", "columns": ["city", "amount"], "rows": [["Tehran", 1000]], "sheet": "شهریور"}).startswith("Created")
    assert "already exists" in tools["write_table"].invoke({"path": "/book.xlsx", "columns": ["a"], "rows": [], "sheet": "مهر"})
    assert "already exists" in tools["write_table"].invoke({"path": "/book.xlsx", "columns": ["a"], "rows": []})

    first = tools["read_table"].invoke({"path": "/book.xlsx"})
    assert "sheets: مهر, شهریور" in first and "rows: 2" in first
    assert "rows: 1" in tools["read_table"].invoke({"path": "/book.xlsx", "sheet": "شهریور"})
    assert "all: 2100.0" in tools["aggregate_table"].invoke({"path": "/book.xlsx", "value_column": "amount"})
    assert "all: 1000.0" in tools["aggregate_table"].invoke({"path": "/book.xlsx", "value_column": "amount", "sheet": "شهریور"})
    assert "no sheet" in tools["read_table"].invoke({"path": "/book.xlsx", "sheet": "آبان"})

    tools["add_table_rows"].invoke({"path": "/book.xlsx", "sheet": "شهریور", "rows": [{"city": "Yazd", "amount": 5}]})
    assert "all: 1005.0" in tools["aggregate_table"].invoke({"path": "/book.xlsx", "value_column": "amount", "sheet": "شهریور"})


# ── archives ──


def test_a_zip_is_listed_and_unpacked_into_a_folder_of_its_own(tools, workspace):
    with zipfile.ZipFile(workspace.root / "docs.zip", "w") as archive:
        archive.writestr("a.txt", "one")
        archive.writestr("sub/b.txt", "two")

    assert "2 files" in tools["list_archive"].invoke({"path": "/docs.zip"})
    assert tools["extract_archive"].invoke({"path": "/docs.zip"}) == "Unpacked 2 files into /docs/."
    assert (workspace.root / "docs" / "sub" / "b.txt").read_text() == "two"
    assert "not empty" in tools["extract_archive"].invoke({"path": "/docs.zip"}), "nothing already there is overwritten"
    assert tools["extract_archive"].invoke({"path": "/docs.zip", "destination": "/.harness/x"}).startswith("Error:")
    assert not any((workspace.harness / "artifacts").glob("unpack-*")), "no half-unpacked folder is left behind"


def test_an_archive_cannot_write_outside_its_folder_or_fill_the_disk(tools, workspace, monkeypatch):
    with zipfile.ZipFile(workspace.root / "slip.zip", "w") as archive:
        archive.writestr("../../escaped.txt", "x")
    with zipfile.ZipFile(workspace.root / "bomb.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("big.bin", b"\0" * 3_000_000)

    assert "unsafe path" in tools["extract_archive"].invoke({"path": "/slip.zip"})
    assert not (workspace.root.parent / "escaped.txt").exists()
    monkeypatch.setattr("harness.tools.archives.MAX_BYTES", 1_000_000)
    assert "more than 1 GB" in tools["extract_archive"].invoke({"path": "/bomb.zip"})
    assert not (workspace.root / "bomb").exists()
    assert tools["list_archive"].invoke({"path": "/notes.txt"}).startswith("Error:")


@pytest.mark.skipif(sys.platform != "win32", reason="uses the bsdtar that ships with Windows")
def test_a_7z_archive_is_read_through_the_system_tar(tools, workspace):
    from harness.tools.archives import _bsdtar

    (workspace.root / "note.txt").write_text("hello", encoding="utf-8")
    subprocess.run([_bsdtar(), "--format=7zip", "-cf", "pack.7z", "note.txt"], cwd=workspace.root, check=True)

    assert "note.txt" in tools["list_archive"].invoke({"path": "/pack.7z"})
    assert tools["extract_archive"].invoke({"path": "/pack.7z", "destination": "/unpacked"}) == "Unpacked 1 files into /unpacked/."
    assert (workspace.root / "unpacked" / "note.txt").read_text(encoding="utf-8") == "hello"


# ── browser ──


@pytest.mark.parametrize("url", ["http://127.0.0.1:8765/files", "http://localhost/", "http://192.168.1.1/", "http://169.254.169.254/latest", "file:///C:/Windows/win.ini", "https://user:pass@example.com/"])
async def test_the_browser_refuses_local_and_non_web_addresses(tools, url):
    assert (await tools["browse_page"].ainvoke({"url": url})).startswith("Error:")


async def test_the_browser_reads_a_page_that_builds_itself_with_script(tools, workspace, monkeypatch):
    page = b"<html><head><title>Price list</title></head><body><h1>Prices</h1><p id=p></p><script>document.getElementById('p').textContent='Saffron ' + (40+2)</script></body></html>"

    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):  # noqa: N802
            self.send_response(200)
            self.send_header("Content-Type", "text/html")
            self.end_headers()
            self.wfile.write(page)

        def log_message(self, *args):
            pass

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()

    async def public(host):  # the test page is local; pretend it is not
        return True

    monkeypatch.setattr(browser, "_is_public", public)
    try:
        result = await tools["browse_page"].ainvoke({"url": f"http://127.0.0.1:{server.server_port}/", "screenshot": True})
    finally:
        server.shutdown()
    if "no browser was found" in result:
        pytest.skip("no Edge or Chrome on this machine")

    assert "title: Price list" in result and "Saffron 42" in result
    shots = list((workspace.harness / "artifacts" / "screenshots").glob("*.png"))
    assert len(shots) == 1 and shots[0].read_bytes().startswith(b"\x89PNG")


# ── the rules cover the new tools ──


def test_new_tools_obey_the_same_rules_as_the_built_in_ones(workspace):
    workspace.rules_file.write_text("files:\n  - {paths: ['/finance/**'], operations: [write], access: deny}\ntools: {execute: ask}\n", encoding="utf-8")
    ask, read_only = (build_policy(load_rules(workspace.rules_file), mode) for mode in ("ask", "read_only"))

    assert ask.decide("write_document", {"path": "/finance/report.docx"}) == "deny"
    assert ask.decide("write_table", {"path": "/finance/new.xlsx"}) == "deny"
    assert ask.decide("extract_archive", {"path": "/in.zip", "destination": "/finance/in"}) == "deny"
    assert ask.decide("extract_archive", {"path": "/finance/in.zip"}) == "deny", "the default destination sits next to the archive"
    assert ask.decide("browse_page", {"url": "https://example.com"}) == "ask", "asks even though these rules never mention the browser"
    assert read_only.decide("write_document", {"path": "/report.docx"}) == "deny"
    assert read_only.decide("browse_page", {"url": "https://example.com"}) == "deny"
    assert read_only.decide("read_document", {"path": "/report.docx"}) == "allow"


def test_a_folder_closed_for_reading_is_closed_to_every_reading_tool(workspace):
    workspace.rules_file.write_text("files:\n  - {paths: ['/secrets/**'], operations: [read], access: deny}\n", encoding="utf-8")
    policy = build_policy(load_rules(workspace.rules_file), "auto")

    for tool in ("read_file", "read_document", "read_table", "aggregate_table", "list_archive"):
        assert policy.decide(tool, {"path": "/secrets/payroll.xlsx", "file_path": "/secrets/payroll.xlsx"}) == "deny", tool
