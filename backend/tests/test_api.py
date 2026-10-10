"""The HTTP layer, exercised end to end against the demo model."""

from __future__ import annotations

import json

import httpx
import pytest

from server import keychain
from server.app import create_app


@pytest.fixture
async def api(tmp_path, monkeypatch):
    """A client on a fresh app, with machine settings and the keychain redirected away from the real ones."""
    monkeypatch.setenv("HARNESS_HOME", str(tmp_path / "home"))
    monkeypatch.delenv("HARNESS_TOKEN", raising=False)
    vault: dict[str, str] = {}
    monkeypatch.setattr(keychain, "get_key", lambda model_id: vault.get(model_id, ""))
    monkeypatch.setattr(keychain, "set_key", vault.__setitem__)
    monkeypatch.setattr(keychain, "delete_key", lambda model_id: vault.pop(model_id, None))

    app = create_app()
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://127.0.0.1:8765") as client:
        client.vault = vault  # type: ignore[attr-defined]
        client.app = app  # type: ignore[attr-defined]
        client.workspace_path = str(tmp_path / "ws")  # type: ignore[attr-defined]
        yield client


async def _open(api) -> None:
    assert (await api.post("/workspace", json={"path": api.workspace_path})).status_code == 200


async def _finish(api, thread_id: str) -> dict:
    await api.app.state.harness.manager.wait(thread_id)
    return (await api.get(f"/runs/{thread_id}")).json()


async def test_nothing_runs_before_a_workspace_is_open(api):
    response = await api.post("/runs", json={"prompt": "سلام"})

    assert response.status_code == 409


async def test_a_task_runs_in_the_background_and_its_trace_is_readable(api):
    await _open(api)
    thread_id = (await api.post("/runs", json={"prompt": "گزارش بنویس", "mode": "auto"})).json()["thread_id"]
    run = await _finish(api, thread_id)

    assert run["status"] == "finished" and run["stop_reason"] == "completed"
    assert run["title"] == "گزارش بنویس"
    assert {"run_started", "plan", "tool_started", "file_changed", "run_finished"} <= {event["type"] for event in run["events"]}
    assert any(file["name"] == "گزارش-نمایشی.md" for file in (await api.get("/files")).json())


async def test_a_task_is_renamed_and_deleted(api, tmp_path):
    await _open(api)
    thread_id = (await api.post("/runs", json={"prompt": "گزارش بنویس", "mode": "auto"})).json()["thread_id"]
    await _finish(api, thread_id)

    assert (await api.put(f"/runs/{thread_id}", json={"title": "  گزارش مهر  "})).json() == {"title": "گزارش مهر"}
    assert (await api.put(f"/runs/{thread_id}", json={"title": "   "})).status_code == 400
    # A follow-up message continues the task and keeps the name the user gave it.
    await api.post("/runs", json={"prompt": "ادامه بده", "thread_id": thread_id, "mode": "auto"})
    assert (await _finish(api, thread_id))["title"] == "گزارش مهر"

    assert (await api.delete(f"/runs/{thread_id}")).json() == {"deleted": True}
    assert (await api.get("/runs")).json() == []
    assert (await api.get(f"/runs/{thread_id}")).status_code == 404
    assert not (tmp_path / "ws" / ".harness" / "runs" / thread_id).exists()
    assert (tmp_path / "ws" / "گزارش-نمایشی.md").exists(), "files the task made stay in the workspace"
    assert (await api.delete(f"/runs/{thread_id}")).status_code == 404
    assert (await api.put("/runs/never-existed", json={"title": "x"})).status_code == 404
    assert not (tmp_path / "ws" / ".harness" / "runs" / "never-existed").exists()


async def test_the_event_stream_replays_a_finished_run(api):
    await _open(api)
    thread_id = (await api.post("/runs", json={"prompt": "گزارش بنویس", "mode": "auto"})).json()["thread_id"]
    await _finish(api, thread_id)

    body = (await api.get(f"/runs/{thread_id}/events")).text
    events = [json.loads(line[6:]) for line in body.splitlines() if line.startswith("data: {\"")]

    assert events[0]["type"] == "run_started" and events[-1]["type"] == "run_finished"
    assert body.rstrip().endswith("event: end\ndata: {}")


async def test_the_approvals_inbox_lists_and_resolves_paused_actions(api):
    await _open(api)
    thread_id = (await api.post("/runs", json={"prompt": "گزارش بنویس", "mode": "ask"})).json()["thread_id"]
    assert (await _finish(api, thread_id))["status"] == "waiting"

    inbox = (await api.get("/approvals")).json()
    assert inbox[0]["thread_id"] == thread_id and inbox[0]["actions"][0]["tool"] == "write_file"

    wrong = await api.post(f"/runs/{thread_id}/decisions", json={"decisions": []})
    assert wrong.status_code == 400

    assert (await api.post(f"/runs/{thread_id}/decisions", json={"decisions": [{"type": "approve"}]})).status_code == 200
    assert (await _finish(api, thread_id))["stop_reason"] == "completed"
    assert (await api.get("/approvals")).json() == []


async def test_model_keys_go_to_the_keychain_and_never_come_back_out(api):
    saved = (await api.post("/models", json={"name": "Claude", "base_url": "https://api.anthropic.com", "model": "claude-opus-5-5", "api_key": "sk-ant-secret"})).json()

    assert api.vault[saved["id"]] == "sk-ant-secret"
    listed = (await api.get("/models")).json()
    assert "sk-ant-secret" not in json.dumps(listed)
    assert next(m for m in listed if m["id"] == saved["id"])["has_key"] is True

    await api.post(f"/models/{saved['id']}/default")
    assert next(m for m in (await api.get("/models")).json() if m["is_default"])["id"] == saved["id"]

    await api.delete(f"/models/{saved['id']}")
    assert saved["id"] not in api.vault


async def test_files_stay_inside_the_workspace_and_system_files_are_opt_in(api):
    await _open(api)
    await api.put("/files/content", json={"path": "/notes/todo.md", "content": "سلام"})

    assert (await api.get("/files/content", params={"path": "/notes/todo.md"})).json()["content"] == "سلام"
    assert (await api.get("/files/content", params={"path": "/../../etc/passwd"})).status_code == 400

    plain = (await api.get("/files")).json()
    with_system = (await api.get("/files", params={"system": True})).json()
    assert not any(file["system"] for file in plain)
    assert any(file["path"] == "/.harness/rules.yaml" for file in with_system)
    assert not any("/runs/" in file["path"] or "state.db" in file["path"] for file in with_system)


async def test_harness_files_are_editable_and_rules_are_validated(api):
    await _open(api)

    assert (await api.get("/skills")).json()[0]["id"] == "weekly-report"
    assert (await api.get("/agents")).json()[0]["id"] == "analyst"

    await api.put("/memory", json={"content": "- مبالغ به تومان"})
    assert (await api.get("/memory")).json()["content"] == "- مبالغ به تومان"

    assert (await api.put("/rules", json={"content": "files: [unclosed"})).status_code == 400
    assert (await api.put("/rules", json={"content": "tools: {execute: sometimes}\n"})).status_code == 400
    assert (await api.put("/rules", json={"content": "files: []\ntools: {}\n"})).status_code == 200


async def test_a_token_is_required_and_only_in_the_header(api, monkeypatch):
    monkeypatch.setenv("HARNESS_TOKEN", "s3cret")

    assert (await api.get("/health")).status_code == 401
    assert (await api.get("/health", params={"token": "s3cret"})).status_code == 401
    assert (await api.get("/health", headers={"Authorization": "Bearer wrong"})).status_code == 401
    assert (await api.get("/health", headers={"Authorization": "Bearer s3cret"})).status_code == 200


async def test_requests_caused_by_other_web_pages_are_refused(api):
    """A site the user visits can make their browser call localhost; it must get nowhere."""
    assert (await api.get("/health", headers={"Origin": "https://evil.example"})).status_code == 403
    assert (await api.get("/health", headers={"Host": "evil.example:8765"})).status_code == 403  # DNS rebinding
    assert (await api.get("/health", headers={"Origin": "app://harness"})).status_code == 200


async def test_a_workspace_must_be_a_folder_of_its_own(api, tmp_path):
    from pathlib import Path

    for too_broad in (str(Path.home()), Path.home().anchor, "relative/folder"):
        assert (await api.post("/workspace", json={"path": too_broad})).status_code == 400
    (tmp_path / "a-file.txt").write_text("x", encoding="utf-8")
    assert (await api.post("/workspace", json={"path": str(tmp_path / "a-file.txt")})).status_code == 400


async def test_a_workspace_has_a_name_that_is_not_its_folder(api):
    assert (await api.post("/workspace", json={"path": api.workspace_path, "name": "فروشگاه نیکان"})).json()["name"] == "فروشگاه نیکان"

    current = (await api.get("/workspace")).json()
    assert current["name"] == "فروشگاه نیکان" and current["known"][0]["name"] == "فروشگاه نیکان"
    assert (await api.put("/workspace/name", json={"name": "  حسابداری  "})).json()["name"] == "حسابداری"
    assert (await api.put("/workspace/name", json={"name": "   "})).status_code == 400
    # Opening it again without a name keeps the one it has.
    assert (await api.post("/workspace", json={"path": api.workspace_path})).json()["name"] == "حسابداری"


async def test_files_made_elsewhere_are_shown_in_the_app(api, tmp_path):
    from docx import Document
    from openpyxl import Workbook

    await _open(api)
    workspace = tmp_path / "ws"
    (workspace / "photo.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"\0" * 20)
    (workspace / "page.html").write_text("<script>alert(1)</script>", encoding="utf-8")
    book = Workbook()
    book.active.title = "مهر"
    book.active.append(["city", "amount"])
    book.active.append(["Tehran", 1200])
    book.create_sheet("شهریور").append(["city", "amount"])
    book.save(workspace / "sales.xlsx")
    document = Document()
    document.add_heading("گزارش", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text, table.cell(1, 0).text, table.cell(1, 1).text = "a", "b", "1", "2"
    document.save(str(workspace / "report.docx"))

    image = await api.get("/files/raw", params={"path": "/photo.png"})
    assert image.status_code == 200 and image.headers["content-type"] == "image/png" and image.content.startswith(b"\x89PNG")
    assert (await api.get("/files/raw", params={"path": "/page.html"})).status_code == 415, "a page is never served from the app's own address"
    assert (await api.get("/files/raw", params={"path": "/../../x.png"})).status_code == 400

    sheet = (await api.get("/files/preview", params={"path": "/sales.xlsx"})).json()
    assert sheet["kind"] == "table" and sheet["sheets"] == ["مهر", "شهریور"] and sheet["columns"] == ["city", "amount"] and sheet["rows"] == [["Tehran", 1200]]
    assert (await api.get("/files/preview", params={"path": "/sales.xlsx", "sheet": "شهریور"})).json()["total_rows"] == 0
    word = (await api.get("/files/preview", params={"path": "/report.docx"})).json()
    assert word["kind"] == "document" and "# گزارش" in word["markdown"] and "| a | b |\n| --- | --- |\n| 1 | 2 |" in word["markdown"]
    assert (await api.get("/files/preview", params={"path": "/photo.png"})).status_code == 415


async def test_sound_and_video_play_through_a_one_file_ticket(api, tmp_path, monkeypatch):
    await _open(api)
    (tmp_path / "ws" / "clip.mp4").write_bytes(bytes(range(256)) * 4)
    (tmp_path / "ws" / "notes.txt").write_text("secret", encoding="utf-8")
    monkeypatch.setenv("HARNESS_TOKEN", "s3cret")
    auth = {"Authorization": "Bearer s3cret"}

    assert (await api.post("/files/media-ticket", json={"path": "/clip.mp4"})).status_code == 401, "getting a ticket needs the token"
    assert (await api.post("/files/media-ticket", json={"path": "/notes.txt"}, headers=auth)).status_code == 415, "a ticket is only for media"
    ticket = (await api.post("/files/media-ticket", json={"path": "/clip.mp4"}, headers=auth)).json()["ticket"]

    whole = await api.get(f"/files/media/{ticket}")  # no token: a player cannot send one
    assert whole.status_code == 200 and whole.headers["content-type"] == "video/mp4" and len(whole.content) == 1024
    piece = await api.get(f"/files/media/{ticket}", headers={"Range": "bytes=10-19"})
    assert piece.status_code == 206 and piece.content == bytes(range(10, 20)), "a player seeks by asking for pieces"
    assert (await api.get("/files/media/not-a-ticket")).status_code == 404
    assert (await api.get("/files/content", params={"path": "/notes.txt"})).status_code == 401, "nothing else opens without the token"


async def test_a_task_id_cannot_point_outside_the_runs_folder(api):
    await _open(api)

    assert (await api.post("/runs", json={"prompt": "سلام", "thread_id": "../../escape"})).status_code == 422
    assert (await api.get("/runs/%2e%2e")).status_code == 422


async def test_a_model_address_must_be_a_web_address(api):
    assert (await api.post("/models", json={"base_url": "file:///etc/passwd", "model": "x"})).status_code == 422


async def test_the_file_list_skips_bulk_folders_and_big_files_do_not_open(api, tmp_path):
    await _open(api)
    workspace = tmp_path / "ws"
    (workspace / "node_modules" / "pkg").mkdir(parents=True)
    (workspace / "node_modules" / "pkg" / "index.js").write_text("x", encoding="utf-8")
    (workspace / "big.txt").write_text("x" * 2_000_001, encoding="utf-8")

    assert not any("node_modules" in file["path"] for file in (await api.get("/files")).json())
    assert (await api.get("/files/content", params={"path": "/big.txt"})).status_code == 413
